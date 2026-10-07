"""Validate the relocated installed artifact with offline OCR and fresh workers."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from openpyxl import Workbook
from semantica.bundle import verify_bundle_inventory
from semantica.source import parse_source, UnsupportedSourceFormatError
from semantica.project_document_quality import assess_document_quality


root = Path(sys.argv[1]).resolve()
destination = Path(sys.argv[2]).resolve()
manifest = json.loads((root / "manifest.json").read_text())
assert Path(sys.executable).is_relative_to(root)
assert "semantica" in str(parse_source.__code__.co_filename)
assert Path(parse_source.__code__.co_filename).is_relative_to(root)
os.environ.update(DOCLING_ARTIFACTS_PATH=str(root / "models"), HF_HUB_OFFLINE="1", TRANSFORMERS_OFFLINE="1", DOCLING_DEVICE="cpu", ORT_DISABLE_TELEMETRY="1")


def invoke(module, request):
    process = subprocess.run([sys.executable, "-I", "-B", "-m", module], input=json.dumps(request) + "\n",
                             text=True, capture_output=True, timeout=180, check=True)
    response = json.loads(process.stdout.splitlines()[-1])
    assert response["ok"], response
    return response["result"]


with tempfile.TemporaryDirectory(prefix="semantica-native-acceptance-") as directory:
    scratch = Path(directory)
    document = Document()
    document.add_paragraph("\u5317\u4eac\u5927\u5b66\u7814\u7a76\u62a5\u544a 73")
    document.save(scratch / "native.docx")
    picture = Image.new("RGB", (1600, 900), "white")
    ImageDraw.Draw(picture).text((100, 160), "\u5317\u4eac\u5927\u5b66\u7814\u7a76\u62a5\u544a 73", fill="black",
        font=ImageFont.truetype(str(root / "models/RapidOcr/fonts/FZYTK.TTF"), 70))
    picture.save(scratch / "scan.pdf", resolution=150)
    picture.save(scratch / "scan.png")
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[5]).shapes.title.text = "Knowledge Evidence 73"
    slides.save(scratch / "slides.pptx")
    workbook = Workbook()
    workbook.active.append(["Knowledge Evidence", 73])
    workbook.save(scratch / "sheet.xlsx")
    (scratch / "page.html").write_text('<html><body><h1>Knowledge Evidence 73</h1></body></html>', encoding="utf-8")
    (scratch / "table.csv").write_text('name,value\nKnowledge Evidence,73\n', encoding="utf-8")
    (scratch / "article.xml").write_text('<?xml version="1.0"?><!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Publishing DTD v1.3 20210610//EN" "JATS-journalpublishing1-3.dtd"><article><front><article-meta><title-group><article-title>Knowledge Evidence 73</article-title></title-group></article-meta></front><body><sec><title>Evidence</title><p>Knowledge Evidence 73</p></sec></body></article>', encoding="utf-8")
    parsed = []
    for name, mime, force, expected in [
        ("native.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", False, "\u5317\u4eac\u5927\u5b66"),
        ("scan.pdf", "application/pdf", True, "\u5317\u4eac\u5927\u5b66"),
        ("scan.png", "image/png", True, "\u5317\u4eac\u5927\u5b66"),
        ("slides.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", False, "Knowledge Evidence"),
        ("sheet.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", False, "Knowledge Evidence"),
        ("page.html", "text/html", False, "Knowledge Evidence"),
        ("table.csv", "text/csv", False, "Knowledge Evidence"),
        ("article.xml", "application/xml", False, "Knowledge Evidence"),
    ]:
        result = parse_source(scratch / name, name=name, mime_type=mime, force_ocr=force)
        assert expected in result.text and "73" in result.text, result.text
        assert result.document["document"] and result.document["doctags"]
        assert result.document["quality"]["status"] == "passed", result.document["quality"]
        if force:
            assert result.origin == "ocr"
        parsed.append({"name": name, "origin": result.origin, "textLength": len(result.text)})
    for name, mime in [("legacy.doc", "application/msword"), ("legacy.wpd", "application/vnd.wordperfect")]:
        (scratch / name).write_bytes(b"legacy")
        try:
            parse_source(scratch / name, name=name, mime_type=mime, force_ocr=False)
        except UnsupportedSourceFormatError as error:
            assert "office-runtime-unavailable" in str(error), error
        else:
            raise AssertionError("legacy input unexpectedly succeeded without Office capability")
    source = scratch / "source.txt"
    source.write_text("Ada Lovelace designed the Analytical Engine.")
    epub = scratch / "book.epub"
    with ZipFile(epub, "w") as archive:
        archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
        archive.writestr("OEBPS/content.opf", '<package><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>')
        archive.writestr("OEBPS/chapter.xhtml", "<html><body>Ada Lovelace designed the Analytical Engine.</body></html>")
    from semantica.source import source_content_revision
    sources = [{"filePath": str(path), "sourceId": f"source-{index}", "name": path.name, "materialRevision": source_content_revision(path), "mimeType": mime}
        for index, (path, mime) in enumerate([(source, "text/plain"), (epub, "application/epub+zip")])]
    parsed_sources = []
    for item in sources:
        result = invoke("semantica.semantic_worker", {"protocol": "semantica.semantic-worker.v1", "id": "native-parse", "method": "parse_source", "params": {
            "source": item, "outputDir": str(scratch / "parsed"), "forceOcr": False, "documentProcessing": {"mode": "local"}}})
        parsed_sources.append({key: result[key] for key in ("sourceId", "artifactPath", "artifactDigest")})
    built = invoke("semantica.semantic_worker", {"protocol": "semantica.semantic-worker.v1", "id": "native-build", "method": "build_semantic_artifacts", "params": {
        "projectId": "native", "baseSnapshot": None, "inputRevision": "sha256:" + "1" * 64, "outputDir": str(scratch / "output"),
        "recipe": {"id": "deterministic", "version": "1", "forceOcrSourceIds": []}, "sources": sources, "parsedSources": parsed_sources, "documentProcessing": {"mode": "local"},
        "release": {key: manifest[key] for key in ("artifactDigest", "schemaDigest", "mediaTypes")},
        "relays": {key: {"authorizationEnv": "OPENAI_API_KEY", "baseUrl": "http://127.0.0.1:9021/v1/" + endpoint, "modelId": "nvidia/llama-nemotron-embed-vl-1b-v2:free", "capability": capability, "receipts": "required",
            **({"contextWindowTokens": 500000, "maxOutputTokens": 393216} if key == "model" else {})}
            for key, endpoint, capability in [("model", "chat/completions", "knowledge.snapshot.generate"), ("embedding", "embeddings", "knowledge.snapshot.embed")]}}})
    artifacts = {item["kind"]: {key: item[key] for key in ("path", "digest", "kind", "mediaType")} for item in built["artifacts"]}
    graph = json.loads(Path(artifacts["semantic-graph"]["path"]).read_text())
    assert graph["protocol"] == "semantica.semantic-graph.v1"
    assert graph["artifact_revision"] == built["semanticGraph"]["artifactRevision"]
    snapshot = {key: value for key, value in graph.items() if key not in
                ("protocol", "artifact_revision", "input_revision", "release_digest", "schema_digest")}
    snapshot.update(protocol="semantica.semantic-artifact.v1", snapshot_id=graph["artifact_revision"],
                    project_id="native", base_snapshot_id=None)
    artifact_path = scratch / "semantic-artifact.json"
    artifact_path.write_text(json.dumps(snapshot, ensure_ascii=False) + "\n")
    import hashlib
    artifact_digest = "sha256:" + hashlib.sha256(artifact_path.read_bytes()).hexdigest()
    queried = invoke("semantica.semantic_query_worker", {"protocol": "semantica.semantic-query.v1", "id": "native-query", "method": "query", "params": {
        "projectId": "native", "snapshotId": graph["artifact_revision"], "snapshot": {
            "path": str(artifact_path), "digest": artifact_digest, "kind": "semantic-graph",
            "mediaType": manifest["mediaTypes"]["semantic-graph"]}, "retrieval": artifacts["retrieval-index"],
        "query": "Ada", "mode": "keyword", "limit": 5}})
    assert queried["contexts"]
    assert all(item["revision"].startswith("sha256:") for item in built["artifacts"] if item["kind"] == "document-representation")
    for representation in snapshot["document_representations"]:
        source = next(item for item in sources if item["sourceId"] == representation["source_id"])
        assert representation["content_hash"] == representation["material_revision_id"] == source["materialRevision"]
    assert assess_document_quality("Broken \ufffd\ufffd mapping", {})["status"] == "needs_review"
    verify_bundle_inventory(root, manifest, remove_untracked_bytecode=True)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps({"artifactDigest": manifest["artifactDigest"], "schemaDigest": manifest["schemaDigest"],
        "parsers": parsed, "queryContexts": len(queried["contexts"]), "documentQualityAdmission": True,
        "representationRevisions": True, "postRunInventoryVerified": True, "deterministicProtocolOnly": True}, indent=2) + "\n")
